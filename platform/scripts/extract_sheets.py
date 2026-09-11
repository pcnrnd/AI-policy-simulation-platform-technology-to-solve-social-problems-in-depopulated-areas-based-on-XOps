"""외부데이터 xlsx → 시트별 CSV 기계 덤프 (platform/data/real/ → platform/data/extracted/).

셀 값을 그대로 옮겨 쓴다: 형 변환·공백 정리·빈 행 제거·병합셀 채움·헤더 판정은
하지 않는다(다음 단계 몫). 이미지·차트는 셀 읽기 경로에 나타나지 않으므로 자동으로
빠지며, MANIFEST 비고에만 존재를 남긴다.

사용법:
    python platform/scripts/extract_sheets.py                       # real/ 4개 파일 전부
    python platform/scripts/extract_sheets.py --input <xlsx> --dataset <slug> --out <dir>
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import logging
from pathlib import Path
from typing import Any

logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
_log = logging.getLogger("extract_sheets")

_REAL_DIR = Path(__file__).resolve().parents[1] / "data" / "real"
_DEFAULT_OUT_DIR = Path(__file__).resolve().parents[1] / "data" / "extracted"

# 기본 모드에서 처리하는 데이터셋 (real/ 파일명과 1:1)
DEFAULT_DATASETS: dict[str, Path] = {
    "namwon_bccard_consumption": _REAL_DIR / "namwon_bccard_consumption_240130.xlsx",
    "gwto_kt_tourism_indicators": _REAL_DIR / "gwto_kt_tourism_indicators_202512.xlsx",
    "namwon_kt_visitors": _REAL_DIR / "namwon_kt_visitors_20231130.xlsx",
    "nowon_kt_festival": _REAL_DIR / "nowon_kt_festival_20250919.xlsx",
}

# (dataset, 원시트명) -> 영문 slug. 매핑에 없는 시트는 sheetNN 으로 대체하고 경고한다.
SHEET_SLUGS: dict[str, dict[str, str]] = {
    "namwon_bccard_consumption": {
        "데이터레이아웃(최종)": "data_layout",
        "1.행정동 업종별": "dong_industry_sales",
    },
    "namwon_kt_visitors": {
        "월별 행정동 방문객수": "monthly_dong_visitors",
        "성연령별 방문객수": "visitors_by_sex_age",
    },
    "nowon_kt_festival": {
        "개요": "overview",
        "일별 방문객수": "daily_visitors",
        "성연령별 방문객수": "visitors_by_sex_age",
        "시간대별 유입 방문객수": "hourly_inflow",
        "시간대별 존재 인구수": "hourly_presence",
        "국적별 방문객수": "visitors_by_nationality",
        "국적별 시간대별 방문객수": "hourly_by_nationality",
        "방문객 거주지역 비율": "residence_ratio",
        "24시간 이후 이동지역 비율": "move_after_24h_ratio",
        "2시간 이후 유출지역 비율": "outflow_after_2h_ratio",
        "방문객 평균 체류시간": "avg_stay_hours",
    },
    "gwto_kt_tourism_indicators": {
        "1-1 외지인 월별": "s1_1_nonlocal_monthly",
        "1-2 전체 관광객 월별": "s1_2_total_tourists_monthly",
        "1-3 관광객 연령": "s1_3_tourists_age",
        "1-4 관광객 성별": "s1_4_tourists_sex",
        "1-5 외국인": "s1_5_foreigners",
        "데이터랩 외국인": "datalab_foreigners",
        "1-6 숙박": "s1_6_lodging",
        "1-7 거주지 광역": "s1_7_residence_region",
        "1-8 거주지 기초 ": "s1_8_residence_city",
        "2-1 월별 소비": "s2_1_consumption_monthly",
        "2-2 소비 연령": "s2_2_consumption_age",
        "2-3 소비 성별": "s2_3_consumption_sex",
        "2-4 소비 업종": "s2_4_consumption_industry",
        "2-5 소비 시간": "s2_5_consumption_hour",
        "2-6 관광객 소비력": "s2_6_tourist_spending_power",
        "3-1 기초 관광객": "s3_1_city_tourists",
        "3-2 기초 연령": "s3_2_city_age",
        "3-3 기초 성별": "s3_3_city_sex",
        "3-4 기초 시간": "s3_4_city_hour",
        "3-5 기초 숙박": "s3_5_city_lodging",
        "3-6 기초 광역": "s3_6_city_region",
        "3-7 기초 소비": "s3_7_city_consumption",
        "4-1 내비데이터 Top 순위(전월)(500위권)": "s4_1_navi_top500_prev_month",
        "4-2 내비데이터 Top 순위(전년)(500위권)": "s4_2_navi_top500_prev_year",
        "4-3 내비데이터 정리본 (500위권)": "s4_3_navi_summary_500",
        "일자별 추이": "daily_trend",
    },
}

# 셀 읽기 경로에 나타나지 않는 이미지·차트 존재를 기록만 하는 비고(시트 인덱스 1에 부착)
_EMBEDDED_OBJECT_NOTES: dict[str, str] = {
    "nowon_kt_festival": "워크북에 이미지 1개(개요 시트), 차트 2개 포함 — 셀 값에는 반영되지 않음",
    "gwto_kt_tourism_indicators": "워크북에 차트 7개 포함 — 셀 값에는 반영되지 않음",
}


def sha256_of(path: Path) -> str:
    """파일 전체 sha256 hex digest."""
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def cell_to_str(value: Any) -> str:
    """None→빈 문자열, datetime 계열→isoformat, 그 외→str(value) 그대로."""
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def slug_for_sheet(dataset: str, sheet_name: str, sheet_index: int) -> str:
    """매핑에 있으면 slug, 없으면 sheetNN + 경고."""
    mapping = SHEET_SLUGS.get(dataset, {})
    slug = mapping.get(sheet_name)
    if slug is None:
        slug = f"sheet{sheet_index:02d}"
        _log.warning(
            "dataset=%s 시트 %r(순번 %d)에 slug 매핑이 없어 %r 을 사용합니다.",
            dataset,
            sheet_name,
            sheet_index,
            slug,
        )
    return slug


def extract_workbook(xlsx_path: Path, dataset: str, out_root: Path) -> list[dict[str, Any]]:
    """워크북의 모든 시트를 CSV로 덤프하고 MANIFEST 행 목록을 반환한다."""
    import openpyxl

    source_sha256 = sha256_of(xlsx_path)
    out_dir = out_root / dataset
    out_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    manifest_rows: list[dict[str, Any]] = []
    embedded_note = _EMBEDDED_OBJECT_NOTES.get(dataset, "")
    try:
        for sheet_index, sheet_name in enumerate(wb.sheetnames, start=1):
            ws = wb[sheet_name]
            slug = slug_for_sheet(dataset, sheet_name, sheet_index)
            csv_path = out_dir / f"{sheet_index:02d}_{slug}.csv"
            rows = ws.max_row or 0
            cols = ws.max_column or 0

            with csv_path.open("w", encoding="utf-8", newline="") as f:
                writer = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
                for row in ws.iter_rows(
                    min_row=1, max_row=rows, min_col=1, max_col=cols, values_only=True
                ):
                    writer.writerow([cell_to_str(v) for v in row])

            manifest_rows.append(
                {
                    "dataset": dataset,
                    "source_file": xlsx_path.name,
                    "source_sha256": source_sha256,
                    "sheet_index": sheet_index,
                    "sheet_name": sheet_name,
                    "csv_path": str(csv_path.relative_to(out_root.parent)).replace("\\", "/"),
                    "rows": rows,
                    "cols": cols,
                    "note": embedded_note if sheet_index == 1 else "",
                }
            )
    finally:
        wb.close()

    return manifest_rows


def write_manifest(manifest_rows: list[dict[str, Any]], out_root: Path) -> Path:
    """MANIFEST.csv 를 덮어쓴다."""
    manifest_path = out_root / "MANIFEST.csv"
    fieldnames = [
        "dataset",
        "source_file",
        "source_sha256",
        "sheet_index",
        "sheet_name",
        "csv_path",
        "rows",
        "cols",
        "note",
    ]
    with manifest_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, quoting=csv.QUOTE_MINIMAL)
        writer.writeheader()
        writer.writerows(manifest_rows)
    return manifest_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, help="처리할 단일 xlsx 경로")
    parser.add_argument("--dataset", type=str, help="--input 사용 시 데이터셋 slug")
    parser.add_argument("--out", type=Path, default=_DEFAULT_OUT_DIR, help="출력 루트 디렉터리")
    args = parser.parse_args()

    if args.input:
        if not args.dataset:
            parser.error("--input 사용 시 --dataset 이 필요합니다.")
        targets = {args.dataset: args.input}
    else:
        targets = DEFAULT_DATASETS

    all_rows: list[dict[str, Any]] = []
    for dataset, xlsx_path in targets.items():
        _log.info("처리 시작: dataset=%s file=%s", dataset, xlsx_path)
        rows = extract_workbook(xlsx_path, dataset, args.out)
        _log.info("처리 완료: dataset=%s 시트 %d개", dataset, len(rows))
        all_rows.extend(rows)

    manifest_path = write_manifest(all_rows, args.out)
    _log.info("MANIFEST 작성: %s (%d행)", manifest_path, len(all_rows))


if __name__ == "__main__":
    main()
