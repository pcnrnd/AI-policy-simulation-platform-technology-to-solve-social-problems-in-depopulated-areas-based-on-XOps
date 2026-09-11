"""extract_sheets.py 검증: 임시 xlsx(빈 셀·숫자·문자열·병합셀 포함)를 덤프해 확인한다."""

from __future__ import annotations

import csv
import importlib.util
import sys
from pathlib import Path

import openpyxl
import pytest

_SCRIPT_PATH = Path(__file__).resolve().parents[1] / "extract_sheets.py"
_spec = importlib.util.spec_from_file_location("extract_sheets", _SCRIPT_PATH)
extract_sheets = importlib.util.module_from_spec(_spec)
sys.modules["extract_sheets"] = extract_sheets
_spec.loader.exec_module(extract_sheets)


@pytest.fixture
def sample_xlsx(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "첫번째 시트"
    ws1.append(["이름", "값", "비고"])
    ws1.append(["가", 1, None])
    ws1.append(["나", 3.14, "텍스트"])
    ws1.merge_cells("A4:B4")
    ws1["A4"] = "병합됨"

    ws2 = wb.create_sheet("두번째 시트")
    ws2.append(["x", "y"])
    ws2.append([10, 20])

    path = tmp_path / "sample.xlsx"
    wb.save(path)
    return path


def test_extract_workbook_dumps_all_sheets(tmp_path: Path, sample_xlsx: Path) -> None:
    out_root = tmp_path / "extracted"
    rows = extract_sheets.extract_workbook(sample_xlsx, "unit_test_ds", out_root)

    assert len(rows) == 2
    out_dir = out_root / "unit_test_ds"
    csv_files = sorted(out_dir.glob("*.csv"))
    assert [f.name for f in csv_files] == ["01_sheet01.csv", "02_sheet02.csv"]


def test_cell_values_are_preserved(tmp_path: Path, sample_xlsx: Path) -> None:
    out_root = tmp_path / "extracted"
    extract_sheets.extract_workbook(sample_xlsx, "unit_test_ds", out_root)

    with (out_root / "unit_test_ds" / "01_sheet01.csv").open(encoding="utf-8", newline="") as f:
        sheet1_rows = list(csv.reader(f))

    assert sheet1_rows[0] == ["이름", "값", "비고"]
    assert sheet1_rows[1] == ["가", "1", ""]
    assert sheet1_rows[2] == ["나", "3.14", "텍스트"]
    # 병합셀: 좌상단만 값을 가지고 나머지 셀은 read_only 워크북에서 None → ""
    assert sheet1_rows[3] == ["병합됨", "", ""]


def test_manifest_rows_match_sheet_dimensions(tmp_path: Path, sample_xlsx: Path) -> None:
    out_root = tmp_path / "extracted"
    rows = extract_sheets.extract_workbook(sample_xlsx, "unit_test_ds", out_root)
    manifest_path = extract_sheets.write_manifest(rows, out_root)

    with manifest_path.open(encoding="utf-8", newline="") as f:
        manifest_rows = list(csv.DictReader(f))

    assert len(manifest_rows) == 2
    sheet1 = next(r for r in manifest_rows if r["sheet_index"] == "1")
    assert sheet1["sheet_name"] == "첫번째 시트"
    assert sheet1["rows"] == "4"
    assert sheet1["cols"] == "3"
    assert sheet1["source_sha256"] == extract_sheets.sha256_of(sample_xlsx)

    sheet2 = next(r for r in manifest_rows if r["sheet_index"] == "2")
    assert sheet2["sheet_name"] == "두번째 시트"
    assert sheet2["rows"] == "2"
    assert sheet2["cols"] == "2"


def test_rerun_is_idempotent(tmp_path: Path, sample_xlsx: Path) -> None:
    out_root = tmp_path / "extracted"
    extract_sheets.extract_workbook(sample_xlsx, "unit_test_ds", out_root)
    first = (out_root / "unit_test_ds" / "01_sheet01.csv").read_bytes()

    extract_sheets.extract_workbook(sample_xlsx, "unit_test_ds", out_root)
    second = (out_root / "unit_test_ds" / "01_sheet01.csv").read_bytes()

    assert first == second
