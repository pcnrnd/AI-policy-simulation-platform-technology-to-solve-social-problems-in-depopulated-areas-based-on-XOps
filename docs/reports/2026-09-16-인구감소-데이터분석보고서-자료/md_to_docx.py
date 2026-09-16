"""Build figures.json, then render the analysis report MD into the PCN template DOCX.

MD stays the content of record. The template supplies cover, header/footer, TOC field and styles;
only the body is replaced. Images are inserted from the same PNG files the MD references.
"""
import hashlib
import json
import re
import shutil
import subprocess
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

KST = timezone(timedelta(hours=9))
REPO = Path('F:/rnd_github/pcnrnd/Population_decline')
MD = REPO / 'docs/reports/2026-09-16-인구감소-데이터분석보고서.md'
ASSET = REPO / 'docs/reports/2026-09-16-인구감소-데이터분석보고서-자료'
DOCX = MD.with_suffix('.docx')
TEMPLATE = Path('C:/Users/PCN/.claude/skills/tech-doc/tech-doc-template.docx')

text = MD.read_text(encoding='utf-8')


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


# ---------------- figures.json ----------------
FIG = [
    {'id': '그림-01', 'title': '원천에서 모델까지의 데이터 계보와 도달 범위', 'kind': 'diagram',
     'source_design': '그림-01.mmd', 'editable_source': '그림-01.pptx', 'editable_slide': 1,
     'evidence': 'stage6-full-data/inventory.json (패밀리·표 수), pg_reader.ALLOWED_TABLES',
     'png': '그림-01.png', 'caption': '그림 1. 원천에서 모델까지의 데이터 계보와 도달 범위'},
    {'id': '차트-01', 'title': '데이터셋별 관측 기간과 관측 단위', 'kind': 'chart',
     'script': 'make_charts.py', 'data': '차트-01.csv', 'png': '차트-01.png',
     'evidence': 'stage6-full-data/quality.json',
     'caption': '그림 2. 데이터셋별 관측 기간과 관측 단위 (빗금 = 원천 공백 2020–2021)'},
    {'id': '차트-05', 'title': '강원 GWTO 지표 테이블별 유효 구간', 'kind': 'chart',
     'script': 'make_charts.py', 'data': '차트-05.csv', 'png': '차트-05.png',
     'evidence': 'stage6-full-data/quality.json (per_table.period)',
     'caption': '그림 3. 강원 GWTO 지표 테이블별 유효 구간 — 시작월이 제각각이라 분모가 다르다'},
    {'id': '차트-04', 'title': '남원 23개 행정동의 점유율 편중', 'kind': 'chart',
     'script': 'make_charts.py', 'data': '차트-04.csv', 'png': '차트-04.png',
     'evidence': 'prepared CSV 직접 집계',
     'caption': '그림 4. 남원 23개 행정동의 점유율 편중 — 각 지표를 자기 합계로 정규화한 값 (단위가 달라 절대량 비교 아님)'},
    {'id': '차트-02', 'title': '남원시 월별 외지인 방문객 합계', 'kind': 'chart',
     'script': 'make_charts.py', 'data': '차트-02.csv', 'png': '차트-02.png',
     'evidence': 'namwon_kt_visitors/01_monthly_dong_visitors.csv',
     'caption': '그림 5. 남원시 월별 외지인 방문객 합계 (23개 행정동 합, 2019-01~2023-10)'},
    {'id': '차트-03', 'title': '남원시 월별 BC카드 추정 매출 합계', 'kind': 'chart',
     'script': 'make_charts.py', 'data': '차트-03.csv', 'png': '차트-03.png',
     'evidence': 'namwon_bccard_consumption/02_dong_industry_sales.csv',
     'caption': '그림 6. 남원시 월별 BC카드 추정 매출 합계 (2019-01~2023-12, 2020–2021 원천 부재)'},
    {'id': '차트-06', 'title': '서울 노원 축제 일별 방문객', 'kind': 'chart',
     'script': 'make_charts.py', 'data': '차트-06.csv', 'png': '차트-06.png',
     'evidence': 'nowon_kt_festival/02_daily_visitors.csv',
     'caption': '그림 7. 서울 노원 축제 일별 방문객 (단일 지점)'},
    {'id': '차트-07', 'title': '구간별 MAE 비', 'kind': 'chart',
     'script': 'make_charts.py', 'data': '차트-07.csv', 'png': '차트-07.png',
     'evidence': 'stage4-model/grid.json — 확정된 기존 증거, 재학습 없음',
     'caption': '그림 8. 구간별 MAE 비 — 확정된 기존 증거 재사용 (재학습 없음)'},
]
for f in FIG:
    f['png_sha256'] = sha(ASSET / f['png'])
    if f.get('data'):
        f['data_sha256'] = sha(ASSET / f['data'])
    if f.get('editable_source'):
        f['editable_source_sha256'] = sha(ASSET / f['editable_source'])
        f['design_sha256'] = sha(ASSET / f['source_design'])
    f['referenced_in_md'] = ('(' + '2026-09-16-인구감소-데이터분석보고서-자료/' + f['png'] + ')') in text
manifest = {
    'built_at_kst': datetime.now(KST).isoformat(timespec='seconds'),
    'commit': subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=REPO, text=True).strip(),
    'rule': 'If the editable source or the underlying data changes, regenerate PNG -> Word -> PDF together '
            'and refresh these hashes.',
    'figures': FIG}
(ASSET / 'figures.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
missing = [f['id'] for f in FIG if not f['referenced_in_md']]
print('figures.json written; figures=%d; not referenced in MD: %s' % (len(FIG), missing or 'none'))

# ---------------- DOCX ----------------
shutil.copy(TEMPLATE, DOCX)
doc = Document(DOCX)


def delete(par):
    par._element.getparent().remove(par._element)


# cover
for p in doc.paragraphs:
    if p.text.strip().startswith('[프로젝트명'):
        p.runs[0].text = '인구감소 R&D 플랫폼'
        for r in p.runs[1:]:
            r.text = ''
    if p.style.name == 'Title':
        for r in p.runs:
            r.text = ''
        p.runs[0].text = '데이터 분석 보고서'
        break

# header text: the template ships "기술문서"; this document is an analysis report.
# Only the literal is replaced - header layout, styles and the logo stay as the template defines them.
for sect in doc.sections:
    for hdr in (sect.header, sect.first_page_header, sect.even_page_header):
        targets = list(hdr.paragraphs)
        for tbl in hdr.tables:                      # 템플릿 머리글은 표 셀 안에 문자열을 둔다
            for row in tbl.rows:
                for cell in row.cells:
                    targets.extend(cell.paragraphs)
        for para in targets:
            for run in para.runs:
                if '기술문서' in run.text:
                    run.text = run.text.replace('기술문서', '데이터 분석 보고서')

# drop guide boxes everywhere
for p in list(doc.paragraphs):
    if p.style.name == 'Guide Box':
        delete(p)

# find the TOC marker, then wipe everything after it
body = doc.element.body
paras = doc.paragraphs
toc_idx = next(i for i, p in enumerate(paras) if p.text.strip().startswith('※ 여기를 클릭'))
keep_until = paras[toc_idx]._element
seen = False
for child in list(body):
    if child is keep_until:
        seen = True
        continue
    if seen and child.tag != keep_until.tag.replace('}p', '}sectPr'):
        if child.tag.endswith('}sectPr'):
            continue
        body.remove(child)

# also drop the mermaid sample block + its caption that sat before the TOC
for p in list(doc.paragraphs):
    if p.style.name in ('Source Code', 'Caption') and (
            'flowchart' in p.text or 'preprocess' in p.text or 'SRLSTM' in p.text
            or 'compute_ade' in p.text or '텐서 [N' in p.text or 'Mermaid 예시' in p.text):
        delete(p)
for p in list(doc.paragraphs):
    if p.text.strip().startswith('공통 규칙. ①'):
        delete(p)

# template document-info tables: fill the two real ones, drop the tool-selection guide table
# 표지는 본문이 선언한 기준 커밋을 그대로 쓴다 — 빌드 시점 HEAD를 쓰면 본문과 어긋난다.
declared = re.search(r'\*\*기준 커밋\*\* `([0-9a-f]{40})`', text)
head = declared.group(1) if declared else subprocess.check_output(
    ['git', 'rev-parse', 'HEAD'], cwd=REPO, text=True).strip()
info = {'작성일': '2026-09-16', '작성자': '[확인 필요]',
        '대상 코드베이스': 'F:/rnd_github/pcnrnd/Population_decline @ ' + head[:7]}
for tbl in list(doc.tables):
    header = ' '.join(c.text.strip() for c in tbl.rows[0].cells)
    if header.startswith('작성일'):
        for row in tbl.rows:
            key = row.cells[0].text.strip()
            if key in info:
                row.cells[1].text = info[key]
            elif key and len(row.cells) > 1 and not row.cells[1].text.strip():
                row.cells[1].text = '[확인 필요]'
    elif header.startswith('버전'):
        if len(tbl.rows) > 1:
            vals = ['1.0', '2026-09-16', '[확인 필요]', '최초 작성']
            for ci, v in enumerate(vals[:len(tbl.rows[1].cells)]):
                tbl.rows[1].cells[ci].text = v
    elif header.startswith('방법'):
        tbl._element.getparent().remove(tbl._element)

INLINE = re.compile(r'(\*\*.+?\*\*|`.+?`)')


CODE = re.compile(r'(`[^`]+`)')


def add_runs(par, s, bold=False):
    """굵게 안에 인라인 코드가 들어간 경우까지 처리한다 — 바깥만 벗기면 백틱이 그대로 남는다."""
    for tok in INLINE.split(s):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            add_runs(par, tok[2:-2], bold=True)
        elif tok.startswith('`') and tok.endswith('`'):
            r = par.add_run(tok[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9)
            r.bold = bold
        else:
            for piece in CODE.split(tok):
                if not piece:
                    continue
                if piece.startswith('`') and piece.endswith('`'):
                    r = par.add_run(piece[1:-1])
                    r.font.name = 'Consolas'
                    r.font.size = Pt(9)
                else:
                    r = par.add_run(piece)
                r.bold = bold


lines = text.split('\n')
i = 0
# skip the H1 + front matter block; it is already on the cover
while i < len(lines) and not lines[i].startswith('## '):
    i += 1

while i < len(lines):
    ln = lines[i]
    s = ln.strip()
    if not s or s == '---':
        i += 1
        continue
    m = re.match(r'^(#{2,4})\s+(.*)$', s)
    if m:
        doc.add_paragraph(m.group(2), style='Heading %d' % (len(m.group(1)) - 1))
        i += 1
        continue
    m = re.match(r'^!\[.*?\]\((.+?)\)$', s)
    if m:
        png = REPO / 'docs/reports' / m.group(1)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(png), width=Cm(16.0))
        i += 1
        while i < len(lines) and not lines[i].strip():   # 이미지와 캡션 사이 빈 줄 건너뛰기
            i += 1
        if i < len(lines) and lines[i].strip().startswith('**그림'):
            cap = doc.add_paragraph(re.sub(r'\*\*', '', lines[i].strip()), style='Caption')
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
        continue
    if s.startswith('|') and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set('|-: '):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not set(''.join(cells)) <= set('-: '):
                rows.append(cells)
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = 'Table Grid'
        for ri, row in enumerate(rows):
            for ci, cell in enumerate(row[:len(rows[0])]):
                c = t.cell(ri, ci)
                c.text = ''
                add_runs(c.paragraphs[0], cell)
                for r in c.paragraphs[0].runs:
                    r.font.size = Pt(9)
                    if ri == 0:
                        r.bold = True
        # 쪽을 넘는 표는 머리행을 반복한다
        hdr = t.rows[0]._tr.get_or_add_trPr()
        el = hdr.makeelement(qn('w:tblHeader'), {})
        hdr.append(el)
        for row in t.rows:
            trPr = row._tr.get_or_add_trPr()
            cant = trPr.makeelement(qn('w:cantSplit'), {})
            trPr.append(cant)
        doc.add_paragraph()
        continue
    if s.startswith('```'):
        i += 1
        while i < len(lines) and not lines[i].strip().startswith('```'):
            doc.add_paragraph(lines[i], style='Source Code')
            i += 1
        i += 1
        continue
    if re.match(r'^[-*]\s+', s):
        add_runs(doc.add_paragraph(style='List Bullet'), re.sub(r'^[-*]\s+', '', s))
        i += 1
        continue
    m = re.match(r'^(\d+)\.\s+(.*)$', s)
    if m:
        # Word의 List Number 스타일은 문서 전체에서 번호를 이어 매긴다 — 절마다 1부터 다시
        # 시작해야 MD와 대응이 유지되므로 번호를 본문 텍스트로 직접 쓴다(독립 검토 지적 반영).
        par = doc.add_paragraph(style='List Paragraph')
        par.add_run(m.group(1) + '. ')
        add_runs(par, m.group(2))
        i += 1
        continue
    if s.startswith('**그림'):
        i += 1
        continue
    add_runs(doc.add_paragraph(), s)
    i += 1

doc.save(DOCX)
d2 = Document(DOCX)
print('docx written: paragraphs=%d tables=%d images=%d' % (
    len(d2.paragraphs), len(d2.tables),
    sum(1 for r in d2.part.rels.values() if 'image' in r.reltype)))
