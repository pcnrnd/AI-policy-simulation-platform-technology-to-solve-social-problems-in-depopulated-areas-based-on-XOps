"""Verify MD <-> DOCX <-> PDF correspondence, figure manifest, and render every PDF page."""
import hashlib
import json
import re
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

import fitz
from docx import Document

KST = timezone(timedelta(hours=9))
REPO = Path('F:/rnd_github/pcnrnd/Population_decline')
RD = REPO / 'docs/reports'
MD = RD / '2026-09-16-인구감소-데이터분석보고서.md'
DOCX = RD / '2026-09-16-인구감소-데이터분석보고서.docx'
PDF = RD / '2026-09-16-인구감소-데이터분석보고서.pdf'
ASSET = RD / '2026-09-16-인구감소-데이터분석보고서-자료'
OUTDIR = Path(sys.argv[1])
PAGES = OUTDIR / 'pdf-pages'
PAGES.mkdir(parents=True, exist_ok=True)


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


md = MD.read_text(encoding='utf-8')
doc = Document(DOCX)
pdf = fitz.open(PDF)

res = {'verified_at_kst': datetime.now(KST).isoformat(timespec='seconds')}

# ---- headings ----
md_heads = [re.sub(r'^#{2,4}\s+', '', l.strip()) for l in md.split('\n') if re.match(r'^#{2,4}\s+', l.strip())]
docx_heads = [p.text.strip() for p in doc.paragraphs
              if p.style.name.startswith('Heading') and p.text.strip() and p.text.strip() != '문서 정보']
pdf_text = '\n'.join(p.get_text() for p in pdf)
res['headings'] = {'md': len(md_heads), 'docx': len(docx_heads),
                   'identical_sequence': md_heads == docx_heads,
                   'missing_in_docx': [h for h in md_heads if h not in docx_heads],
                   'all_present_in_pdf': all(h.split(' ')[0] in pdf_text for h in md_heads)}

# ---- tables ----
md_tables = []
lines = md.split('\n')
i = 0
while i < len(lines):
    if lines[i].strip().startswith('|') and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set('|-: '):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not set(''.join(cells)) <= set('-: '):
                rows.append(cells)
            i += 1
        md_tables.append(rows)
    else:
        i += 1
docx_tables = [t for t in doc.tables if not (t.rows and t.rows[0].cells[0].text.strip() in ('작성일', '버전'))]
cell_mismatch = []
for idx, (mt, dt) in enumerate(zip(md_tables, docx_tables)):
    if len(mt) != len(dt.rows) or len(mt[0]) != len(dt.columns):
        cell_mismatch.append({'table': idx, 'md_shape': [len(mt), len(mt[0])],
                              'docx_shape': [len(dt.rows), len(dt.columns)]})
        continue
    for r, row in enumerate(mt):
        for c, cell in enumerate(row):
            plain = re.sub(r'\*\*|`', '', cell)
            got = dt.cell(r, c).text.strip()
            if plain != got:
                cell_mismatch.append({'table': idx, 'row': r, 'col': c, 'md': plain[:40], 'docx': got[:40]})
res['tables'] = {'md': len(md_tables), 'docx_body': len(docx_tables),
                 'count_match': len(md_tables) == len(docx_tables),
                 'cell_mismatches': cell_mismatch[:10], 'cell_mismatch_count': len(cell_mismatch)}

# ---- figures ----
man = json.loads((ASSET / 'figures.json').read_text(encoding='utf-8'))
md_imgs = re.findall(r'!\[.*?\]\((.+?)\)', md)
docx_imgs = sum(1 for r in doc.part.rels.values() if 'image' in r.reltype)
pdf_imgs = sum(len(p.get_images(full=True)) for p in pdf)
fig_checks = []
for f in man['figures']:
    live = sha(ASSET / f['png'])
    fig_checks.append({'id': f['id'], 'png_hash_matches_manifest': live == f['png_sha256'],
                       'referenced_in_md': f['referenced_in_md'],
                       'caption_in_docx': any(f['caption'] in p.text for p in doc.paragraphs),
                       'caption_in_pdf': f['caption'].split('—')[0].strip()[:24] in pdf_text})
res['figures'] = {'manifest_count': len(man['figures']), 'md_image_refs': len(md_imgs),
                  'docx_embedded_images': docx_imgs, 'pdf_embedded_images': pdf_imgs,
                  'per_figure': fig_checks,
                  'all_hashes_match': all(c['png_hash_matches_manifest'] for c in fig_checks),
                  'all_captions_in_docx': all(c['caption_in_docx'] for c in fig_checks)}

# ---- render every page ----
page_notes = []
for n, page in enumerate(pdf, 1):
    pix = page.get_pixmap(dpi=110)
    out = PAGES / ('page-%02d.png' % n)
    pix.save(out)
    txt = page.get_text().strip()
    page_notes.append({'page': n, 'chars': len(txt), 'images': len(page.get_images(full=True)),
                       'blank': len(txt) == 0 and len(page.get_images(full=True)) == 0,
                       'png': out.name})
res['render'] = {'pages': len(pdf), 'blank_pages': [p['page'] for p in page_notes if p['blank']],
                 'pages_with_images': [p['page'] for p in page_notes if p['images']],
                 'per_page': page_notes, 'page_png_dir': str(PAGES)}

# ---- final hashes ----
res['final_hashes'] = {'md': sha(MD), 'docx': sha(DOCX), 'pdf': sha(PDF),
                       'figures_json': sha(ASSET / 'figures.json')}

(OUTDIR / 'report-verification.json').write_text(json.dumps(res, ensure_ascii=False, indent=2) + '\n',
                                                 encoding='utf-8')
print('headings md/docx %d/%d identical=%s' % (res['headings']['md'], res['headings']['docx'],
                                               res['headings']['identical_sequence']))
print('tables  md/docx %d/%d mismatched cells=%d' % (res['tables']['md'], res['tables']['docx_body'],
                                                     res['tables']['cell_mismatch_count']))
print('figures %d | md refs %d | docx imgs %d | pdf imgs %d | hashes match=%s | captions in docx=%s'
      % (res['figures']['manifest_count'], res['figures']['md_image_refs'],
         res['figures']['docx_embedded_images'], res['figures']['pdf_embedded_images'],
         res['figures']['all_hashes_match'], res['figures']['all_captions_in_docx']))
print('pages %d | blank %s | pages with images %s' % (res['render']['pages'], res['render']['blank_pages'],
                                                     res['render']['pages_with_images']))
for k, v in res['final_hashes'].items():
    print('  %-12s %s' % (k, v))
